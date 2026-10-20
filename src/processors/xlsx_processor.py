import os
import shutil
import time
import re
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles.colors import Color
from concurrent.futures import ThreadPoolExecutor, as_completed

from core.translation_api import call_gemini_batch, call_gemini_fallback, build_prompt
from core.text_utils import load_glossary_df, get_matching_glossary_rules

def cell_to_tagged_text(cell):
    """Convertit une cellule Excel en texte avec balises HTML si elle contient du Rich Text."""
    val = cell.value
    if val is None:
        return ""
        
    if type(val).__name__ == 'CellRichText':
        tagged_str = ""
        for element in val:
            if isinstance(element, str):
                tagged_str += element
            else: # C'est un TextBlock (texte avec style spécifique)
                txt = element.text
                f = element.font
                if not txt: continue
                
                is_b = f.b
                is_i = f.i
                is_u = f.u
                is_s = f.strike
                
                # Extraction sécurisée de la couleur RGB
                c_rgb = None
                if f.color:
                    try:
                        if hasattr(f.color, 'rgb') and isinstance(f.color.rgb, str) and f.color.rgb != "00000000":
                            c_rgb = f.color.rgb
                    except Exception:
                        pass
                
                # Application des balises autour du segment de texte
                if c_rgb: txt = f'<c c="{c_rgb}">{txt}</c>'
                if is_s: txt = f"<s>{txt}</s>"
                if is_u: txt = f"<u>{txt}</u>"
                if is_i: txt = f"<i>{txt}</i>"
                if is_b: txt = f"<b>{txt}</b>"
                
                tagged_str += txt
        return tagged_str
        
    elif isinstance(val, str):
        return val
        
    return str(val)

def tagged_text_to_rich_text(tagged_text):
    """Reconvertit un texte balisé en objet CellRichText compréhensible par Excel."""
    # Si aucune balise n'est trouvée, on renvoie un texte simple (préserve le style global de la cellule)
    if not re.search(r'<(/?[bius]|c c="[^"]+")>', tagged_text):
        return tagged_text 
        
    parts = re.split(r'(</?(?:b|i|u|s)>|<c c="[^"]+">|</c>)', tagged_text)
    rich_text = CellRichText()
    
    curr_b, curr_i, curr_u, curr_s, curr_c = False, False, False, False, None
    
    for part in parts:
        if not part: continue
        
        if part == '<b>': curr_b = True
        elif part == '</b>': curr_b = False
        elif part == '<i>': curr_i = True
        elif part == '</i>': curr_i = False
        elif part == '<u>': curr_u = True
        elif part == '</u>': curr_u = False
        elif part == '<s>': curr_s = True
        elif part == '</s>': curr_s = False
        elif part.startswith('<c c="'):
            m = re.match(r'<c c="([^"]+)">', part)
            if m: curr_c = m.group(1)
        elif part == '</c>':
            curr_c = None
        else:
            # Construction du bloc de texte avec son style en cours
            font = InlineFont(b=curr_b, i=curr_i, u='single' if curr_u else None, strike=curr_s)
            if curr_c:
                font.color = Color(rgb=curr_c)
            rich_text.append(TextBlock(font, part))
            
    return rich_text if len(rich_text) > 0 else tagged_text

def extract_texts_from_excel(wb):
    """Extrait les textes et le balisage des cellules."""
    cells_data = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val and (isinstance(val, str) or type(val).__name__ == 'CellRichText'):
                    # Ignorer les formules
                    if isinstance(val, str) and val.startswith('='):
                        continue
                    
                    tagged = cell_to_tagged_text(cell).strip()
                    if tagged:
                        cells_data.append((ws.title, cell.row, cell.column, tagged))
    return cells_data

def _process_single_batch(args):
    """Traite un lot Excel et renvoie les traductions."""
    batch_idx, batch_cells, glossary_snippet, direction = args
    orig_texts = [cell[3] for cell in batch_cells]
    
    extra_rules = (
        "\n\nCONSIGNES DE STYLE EXCEL ET BALISAGE :\n"
        "1. PRIORITÉ ABSOLUE AU GLOSSAIRE : Applique strictement les traductions du glossaire.\n"
        "2. CONSERVE TOUTES LES BALISES STRUCTURALES : <b>, <i>, <u>, <s> et <c c=\"...\">.\n"
        "3. MOTS ISOLÉS COLORÉS : Si une phrase a une balise de couleur (ex: <c c=\"FFFFC000\">texte</c>), entoure son équivalent traduit avec EXACTEMENT la même balise.\n"
        "4. CONCISION : Les textes de cellules doivent rester courts."
    )
    combined_context = (glossary_snippet or "") + extra_rules

    prompt = build_prompt(orig_texts, combined_context, direction=direction, is_batch=True)
    
    batch_in_tokens = 0
    batch_out_tokens = 0

    translations, in_tok, out_tok = call_gemini_batch(prompt, expected_count=len(batch_cells))
    batch_in_tokens += in_tok
    batch_out_tokens += out_tok

    if translations:
        return batch_idx, batch_cells, translations, orig_texts, batch_in_tokens, batch_out_tokens

    # Mode secours unitaire
    print(f"⚠️ Lot Excel {batch_idx + 1} : Bascule en mode de secours unitaire...")
    fallback_translated = []
    
    for orig_txt in orig_texts:
        fb_prompt = build_prompt(orig_txt, combined_context, direction=direction, is_batch=False)
        fb_trans, fb_in, fb_out = call_gemini_fallback(fb_prompt)
        
        batch_in_tokens += fb_in
        batch_out_tokens += fb_out
        
        if fb_trans:
            fallback_translated.append(fb_trans)
        else:
            fallback_translated.append(orig_txt)
            
    return batch_idx, batch_cells, fallback_translated, orig_texts, batch_in_tokens, batch_out_tokens

def translate_xlsx_universal(input_path: str, output_path: str, glossary_path: str = None, progress_callback=None, direction="Allemand ➔ Français"):
    start_time = time.time()
    total_input_tokens = 0
    total_output_tokens = 0

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    tmp_in = f"/tmp/in_{os.path.basename(input_path)}"
    tmp_out = f"/tmp/out_{os.path.basename(input_path)}"
    
    shutil.copy(input_path, tmp_in)
    
    # Chargement du classeur en demandant explicitement de lire le Rich Text
    keep_vba = tmp_in.lower().endswith('.xlsm')
    try:
        wb = openpyxl.load_workbook(tmp_in, keep_vba=keep_vba, rich_text=True)
    except TypeError:
        # Fallback pour les anciennes versions d'openpyxl
        wb = openpyxl.load_workbook(tmp_in, keep_vba=keep_vba)
        
    glossary_df = load_glossary_df(glossary_path)

    cells_to_translate = extract_texts_from_excel(wb)
    total_cells = len(cells_to_translate)
    
    batch_size, max_workers = (40, 6) if total_cells > 200 else ((20, 3) if total_cells > 40 else (10, 1))

    batch_args_list = [] 
    for i in range(0, total_cells, batch_size):
        batch = cells_to_translate[i:i + batch_size]
        batch_text = " ".join([cell[3] for cell in batch])
        glossary_snippet = get_matching_glossary_rules(batch_text, glossary_df, direction=direction)
        
        batch_args_list.append((i//batch_size, batch, glossary_snippet, direction))

    if not batch_args_list:
        wb.save(tmp_out)
        shutil.copy(tmp_out, output_path)
        os.remove(tmp_in)
        os.remove(tmp_out)
        return {
            "status": "success", "output_path": output_path, "time_seconds": round(time.time() - start_time, 2),
            "input_tokens": 0, "output_tokens": 0, "total_cost_usd": 0.0
        }

    print(f"⚡ Traitement XLSX de {total_cells} cellules...")
    total_batches = len(batch_args_list)
    completed_batches = 0

    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        for future in as_completed([executor.submit(_process_single_batch, arg) for arg in batch_args_list]):
            idx, batch, translated_texts, orig_texts, batch_in, batch_out = future.result()
            
            total_input_tokens += batch_in
            total_output_tokens += batch_out
            
            # Injection des traductions dans le classeur
            for cell_info, t_text in zip(batch, translated_texts):
                ws_title, row_idx, col_idx, _ = cell_info
                ws = wb[ws_title]
                cell = ws.cell(row=row_idx, column=col_idx)
                
                # Reconversion du balisage texte en objet stylisé natif
                cell.value = tagged_text_to_rich_text(t_text)
                
            completed_batches += 1
            if progress_callback:
                progress_callback(completed_batches, total_batches)

    wb.save(tmp_out)
    shutil.copy(tmp_out, output_path)
    os.remove(tmp_in)
    os.remove(tmp_out)
    
    execution_time = time.time() - start_time
    total_cost = ((total_input_tokens / 1_000_000) * 0.075) + ((total_output_tokens / 1_000_000) * 0.30)

    return {
        "status": "success", "output_path": output_path, "time_seconds": round(execution_time, 2),
        "input_tokens": total_input_tokens, "output_tokens": total_output_tokens, "total_cost_usd": total_cost
    }