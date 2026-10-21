"""
Traduction PPTX en gardant la forme et le style.
"""
import copy
import math
import os
import re
import shutil
import tempfile
import time

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from core.segments import Segment
from core.text_fit import Para, Run, get_font, layout
from core.text_utils import load_glossary_df
from core.translation_api import (LANG, PRICE_IN, PRICE_OUT, items_to_tagged,
                                  shorten_segments, translate_segments)

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
A = "{%s}" % A_NS
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
EMU_PER_PT = 12700

LOCALE = {"fr": "fr-FR", "de": "de-DE"}

# Réduction de la taille du texte
SHORTEN_BELOW = 0.85
MIN_SCALE = 0.6
SCALE_STEP = 0.025

_NEUTRAL_RPR_ATTRS = {"lang", "altLang", "dirty", "err", "noProof", "smtClean", "smtId", "bmk"}

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
_OTHER_TYPES = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER,
                PP_PLACEHOLDER.HEADER}

# --------------------------------------------------------------------------
# Segments : paragraphe DrawingML <-> texte balise
# --------------------------------------------------------------------------

def _style_key(rpr):
    if rpr is None:
        return "-"
    el = copy.deepcopy(rpr)
    for a in list(el.attrib):
        if a in _NEUTRAL_RPR_ATTRS:
            del el.attrib[a]
    return etree.tostring(el, encoding="unicode")

_CONNECTORS = set("""
et ou de des du la le les l en à au aux pour par sur avec sans dans d un une que qui
und oder der die das den dem des im in mit für von zu zur zum bei auf an aus nach ein eine
and or of the to for with
""".split())

def _para_text(p):
    return "".join(t.text or "" for t in p.iter(A + "t"))

def _has_bullet(p, chain):
    ppr = p.find(A + "pPr")
    lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
    cands = [ppr] + [Chain._lvl(el, lvl) for el in chain.lst]
    for c in cands:
        if c is None:
            continue
        if c.find(A + "buNone") is not None:
            return False
        if c.find(A + "buChar") is not None or c.find(A + "buAutoNum") is not None or c.find(A + "buBlip") is not None:
            return True
    return False

def is_continuation(prev, nxt, body):
    a, b = _para_text(prev).rstrip(), _para_text(nxt).lstrip()
    if not a or not b:
        return False

    # Protection des champs dynamiques (numéros de diapositive) sans erreur d'attribut
    el_prev = getattr(prev, '_element', prev)
    el_nxt = getattr(nxt, '_element', nxt)
    if el_prev.find(A + "fld") is not None or el_nxt.find(A + "fld") is not None:
        return False

    if a[-1] in ".:;!?),•" or b[0] in "-•–*·>" or re.match(r"^\d+[.)]", b):
        return False
    pp, pn = prev.find(A + "pPr"), nxt.find(A + "pPr")
    if (pp.get("lvl", "0") if pp is not None else "0") != (pn.get("lvl", "0") if pn is not None else "0"):
        return False
    if _has_bullet(nxt, body.chain):
        return False
    words = re.findall(r"[^\W\d_]+", a.lower())
    linked = bool(words) and words[-1] in _CONNECTORS or a[-1] in "-/&+"
    fill = None
    if body.width:
        try:
            l, r, _, _ = _insets(body)
            para = measure_paragraph(body.chain, prev, original_runs(prev))
            avail = body.width - l - r - para.indent
            _, _, w = layout([para], 1e9, 1.0, 0.0, False)
            if avail > 0:
                fill = 1.0 if w >= avail else w / avail
        except Exception:
            fill = None
    if fill is not None and fill < 0.5:
        return False
    if b[0].islower() or linked:
        return True
    return fill is not None and fill >= 0.85 and b[0].isalpha()

def group_segment(paras):
    seg = paragraph_segment(paras[0], require_text=False)
    for nxt in paras[1:]:
        prev_text = "".join(it[2] for it in seg.items if it[0] == "t")
        nseg = paragraph_segment(nxt, require_text=False)
        if prev_text.endswith(("-", "/")) and not prev_text.endswith((" -", " /")):
            pass
        elif prev_text and not prev_text.endswith(" "):
            last_style = next((it for it in reversed(seg.items) if it[0] == "t"), None)
            sid = last_style[1] if last_style else 0
            seg.items.append(("t", sid, " "))
        for it in nseg.items:
            if it[0] == "t":
                key = nseg.meta["keys"][it[1]]
                seg.add_text(it[2], key, nseg.styles[it[1]])
            elif it[0] == "x":
                seg.add_object(nseg.objects[it[1]])
            else:
                seg.add_break(it[0])
            if it[0] == "br":
                seg.meta.setdefault("br", nseg.meta.get("br"))
    seg.meta["merged"] = list(paras[1:])
    return seg if seg.has_text() else None

def paragraph_segment(p, require_text=True):
    seg = Segment()
    seg.meta["p"] = p
    seg.meta["keys"] = {}
    for child in p:
        tag = child.tag
        if tag == A + "pPr" or tag == A + "endParaRPr":
            continue
        if tag == A + "r":
            rpr = child.find(A + "rPr")
            t = child.find(A + "t")
            text = t.text if t is not None and t.text else ""
            parts = text.split("\t")
            key = _style_key(rpr)
            for k, part in enumerate(parts):
                if k:
                    seg.add_break("tab")
                seg.add_text(part, key, rpr)
                if part:
                    seg.meta["keys"][seg.items[-1][1]] = key
        elif tag == A + "br":
            seg.meta.setdefault("br", child)
            seg.add_break("br")
        else:
            seg.add_object(child)
    if not require_text:
        return seg
    return seg if seg.has_text() else None

def rebuild_paragraph(seg, items, target_lang):
    p = seg.meta["p"]
    end = p.find(A + "endParaRPr")
    for child in list(p):
        if child.tag not in (A + "pPr", A + "endParaRPr"):
            p.remove(child)

    def insert(el):
        if end is not None:
            end.addprevious(el)
        else:
            p.append(el)

    locale = LOCALE.get(target_lang)

    def new_run(text, rpr):
        r = etree.Element(A + "r")
        if rpr is not None:
            rp = copy.deepcopy(rpr)
            if locale and rp.get("lang"):
                rp.set("lang", locale)
            r.append(rp)
        t = etree.SubElement(r, A + "t")
        t.text = text
        insert(r)

    def new_br(rpr):
        src = seg.meta.get("br")
        if src is not None:
            br = copy.deepcopy(src)
        else:
            br = etree.Element(A + "br")
            if rpr is not None:
                br.append(copy.deepcopy(rpr))
        insert(br)

    last_rpr = seg.styles.get(seg.base_style())
    for it in items:
        kind = it[0]
        if kind == "t":
            rpr = seg.styles.get(it[1])
            for k, piece in enumerate(it[2].split("\n")):
                if k:
                    new_br(rpr)
                if piece:
                    new_run(piece, rpr)
            last_rpr = rpr
        elif kind == "tab":
            new_run("\t", last_rpr)
        elif kind == "br":
            new_br(last_rpr)
        elif kind == "x":
            insert(seg.objects[it[1]])
    if end is not None and locale and end.get("lang"):
        end.set("lang", locale)
    for other in seg.meta.get("merged", []):
        parent = other.getparent()
        if parent is not None:
            parent.remove(other)

# --------------------------------------------------------------------------
# Styles effectifs (héritage PowerPoint)
# --------------------------------------------------------------------------

class Theme:
    def __init__(self, master=None):
        self.major = self.minor = "Arial"
        if master is None:
            return
        try:
            part = master.part.part_related_by(RT.THEME)
            root = etree.fromstring(part.blob)
            mj = root.find(".//" + A + "majorFont/" + A + "latin")
            mn = root.find(".//" + A + "minorFont/" + A + "latin")
            if mj is not None:
                self.major = mj.get("typeface") or self.major
            if mn is not None:
                self.minor = mn.get("typeface") or self.minor
        except Exception:
            pass

    def resolve(self, face):
        if not face:
            return None
        if face.startswith("+mj"):
            return self.major
        if face.startswith("+mn"):
            return self.minor
        return face

class Chain:
    def __init__(self, lst_styles, body_prs, theme, title=False):
        self.lst = [x for x in lst_styles if x is not None]
        self.body = [x for x in body_prs if x is not None]
        self.theme = theme or Theme()
        self.title = title

    @staticmethod
    def _lvl(el, lvl):
        return el.find(A + "lvl%dpPr" % (lvl + 1))

    def ppr_attr(self, p, lvl, name):
        ppr = p.find(A + "pPr")
        if ppr is not None and ppr.get(name) is not None:
            return ppr.get(name)
        for el in self.lst:
            lp = self._lvl(el, lvl)
            if lp is not None and lp.get(name) is not None:
                return lp.get(name)
        return None

    def ppr_spacing(self, p, lvl, tag):
        cands = []
        ppr = p.find(A + "pPr")
        if ppr is not None:
            cands.append(ppr)
        cands += [self._lvl(el, lvl) for el in self.lst]
        for c in cands:
            if c is None:
                continue
            sp = c.find(A + tag)
            if sp is not None and len(sp):
                v = sp[0]
                if v.tag == A + "spcPct":
                    return ("pct", int(v.get("val", "100000")) / 100000.0)
                if v.tag == A + "spcPts":
                    return ("pts", int(v.get("val", "0")) / 100.0)
        return None

    def rpr_attr(self, rpr, lvl, name):
        if rpr is not None and rpr.get(name) is not None:
            return rpr.get(name)
        for el in self.lst:
            lp = self._lvl(el, lvl)
            if lp is not None:
                d = lp.find(A + "defRPr")
                if d is not None and d.get(name) is not None:
                    return d.get(name)
        return None

    def latin(self, rpr, lvl):
        if rpr is not None:
            lat = rpr.find(A + "latin")
            if lat is not None and lat.get("typeface"):
                return self.theme.resolve(lat.get("typeface"))
        for el in self.lst:
            lp = self._lvl(el, lvl)
            if lp is not None:
                lat = lp.find(A + "defRPr/" + A + "latin")
                if lat is not None and lat.get("typeface"):
                    return self.theme.resolve(lat.get("typeface"))
        return self.theme.major if self.title else self.theme.minor

    def body_attr(self, name, default=None):
        for b in self.body:
            if b.get(name) is not None:
                return b.get(name)
        return default

    def autofit(self):
        for b in self.body:
            for tag in ("normAutofit", "spAutoFit", "noAutofit"):
                el = b.find(A + tag)
                if el is not None:
                    return tag, el
        return None, None

def _bool(v):
    return v in ("1", "true", "on")

def run_metrics(chain, rpr, lvl):
    size = chain.rpr_attr(rpr, lvl, "sz")
    size = int(size) / 100.0 if size else 18.0
    bold = _bool(chain.rpr_attr(rpr, lvl, "b"))
    italic = _bool(chain.rpr_attr(rpr, lvl, "i"))
    return size, get_font(chain.latin(rpr, lvl), bold, italic)

def measure_paragraph(chain, p, runs):
    ppr = p.find(A + "pPr")
    lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
    out = []
    for text, rpr in runs:
        size, font = run_metrics(chain, rpr, lvl)
        out.append(Run(text, size, font))
    if not out:
        size, font = run_metrics(chain, p.find(A + "endParaRPr"), lvl)
        out.append(Run("", size, font))
    ls = chain.ppr_spacing(p, lvl, "lnSpc") or ("pct", 1.0)
    first_size = out[0].size

    def spc(tag):
        v = chain.ppr_spacing(p, lvl, tag)
        if not v:
            return 0.0
        return v[1] if v[0] == "pts" else v[1] * first_size * 1.2

    mar = chain.ppr_attr(p, lvl, "marL")
    ind = chain.ppr_attr(p, lvl, "indent")
    return Para(out, ls, spc("spcBef"), spc("spcAft"),
                indent=(int(mar) / EMU_PER_PT) if mar else 0.0,
                first_indent=(int(ind) / EMU_PER_PT) if ind and int(ind) > 0 else 0.0)

def original_runs(p):
    runs = []
    for child in p:
        if child.tag in (A + "r", A + "fld"):
            t = child.find(A + "t")
            runs.append(((t.text or "") if t is not None else "", child.find(A + "rPr")))
        elif child.tag == A + "br":
            runs.append(("\n", child.find(A + "rPr")))
    return runs

def translated_runs(seg, items):
    runs = []
    last = seg.styles.get(seg.base_style())
    for it in items:
        if it[0] == "t":
            runs.append((it[2], seg.styles.get(it[1])))
            last = seg.styles.get(it[1])
        elif it[0] == "br":
            runs.append(("\n", last))
        elif it[0] == "tab":
            runs.append(("    ", last))
        elif it[0] == "x":
            el = seg.objects[it[1]]
            t = el.find(A + "t")
            if t is not None and t.text:
                runs.append((t.text, el.find(A + "rPr")))
    return runs

# --------------------------------------------------------------------------
# Inventaire des corps de texte
# --------------------------------------------------------------------------

class Body:
    def __init__(self, txBody, chain, width=None, height=None, kind="shape",
                 context="", shape=None, slide_h=None, rects=None, cell=None,
                 rect=None, framed=False):
        self.txBody = txBody
        self.chain = chain
        self.width = width
        self.height = height
        self.kind = kind
        self.context = context
        self.shape = shape
        self.slide_h = slide_h
        self.rects = rects if rects is not None else []
        self.rect = rect
        self.framed = framed
        self.cell = cell
        self.paras = txBody.findall(A + "p")
        self.seg_index = {}
        self.merged = set()

def _insets(body):
    if body.kind == "cell":
        tc = body.cell[3].find(A + "tcPr")
        def get(k, d):
            return (int(tc.get(k)) if tc is not None and tc.get(k) else d) / EMU_PER_PT
        return get("marL", 91440), get("marR", 91440), get("marT", 45720), get("marB", 45720)

    def g(k, d):
        return int(body.chain.body_attr(k, d)) / EMU_PER_PT
    return g("lIns", 91440), g("rIns", 91440), g("tIns", 45720), g("bIns", 45720)

class Inventory:
    def __init__(self, prs):
        self.prs = prs
        self.bodies = []
        self.blob_parts = []
        self.default_style = prs.part._element.find(P + "defaultTextStyle")
        self.slide_h = prs.slide_height / EMU_PER_PT if prs.slide_height else None
        self._themes = {}

    def theme(self, master):
        key = id(master.part)
        if key not in self._themes:
            self._themes[key] = Theme(master)
        return self._themes[key]

    def chain_for_shape(self, shape, master):
        tb = shape._element.find(".//" + P + "txBody")
        lst, bodies = [], []
        if tb is not None:
            lst.append(tb.find(A + "lstStyle"))
            bodies.append(tb.find(A + "bodyPr"))
        title = False
        tx = master._element.find(P + "txStyles")
        if getattr(shape, "is_placeholder", False):
            try:
                ph_type = shape.placeholder_format.type
            except Exception:
                ph_type = None
            title = ph_type in _TITLE_TYPES
            base = shape
            for _ in range(2):
                try:
                    base = base._base_placeholder
                except Exception:
                    base = None
                if base is None:
                    break
                btb = base._element.find(".//" + P + "txBody")
                if btb is not None:
                    lst.append(btb.find(A + "lstStyle"))
                    bodies.append(btb.find(A + "bodyPr"))
            if tx is not None:
                name = "titleStyle" if title else ("otherStyle" if ph_type in _OTHER_TYPES else "bodyStyle")
                lst.append(tx.find(P + name))
        elif tx is not None:
            lst.append(tx.find(P + "otherStyle"))
        lst.append(self.default_style)
        return Chain(lst, bodies, self.theme(master), title)

    def walk_shapes(self, shapes, master, where, xf=(0, 0, 0, 0, 1.0, 1.0), rects=None):
        rects = rects if rects is not None else []
        ox, oy, chx, chy, sx, sy = xf
        for shape in shapes:
            try:
                st = shape.shape_type
            except Exception:
                st = None
            try:
                x0 = ox + (shape.left - chx) * sx
                y0 = oy + (shape.top - chy) * sy
                w_emu, h_emu = shape.width * sx, shape.height * sy
                rect = (x0 / EMU_PER_PT, y0 / EMU_PER_PT,
                        (x0 + w_emu) / EMU_PER_PT, (y0 + h_emu) / EMU_PER_PT)
            except Exception:
                rect = None
            if st == MSO_SHAPE_TYPE.GROUP:
                try:
                    xfrm = shape._element.grpSpPr.find(A + "xfrm")
                    off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
                    choff, chext = xfrm.find(A + "chOff"), xfrm.find(A + "chExt")
                    gx = ox + (int(off.get("x")) - chx) * sx
                    gy = oy + (int(off.get("y")) - chy) * sy
                    cw, chh = int(chext.get("cx")), int(chext.get("cy"))
                    gsx = sx * (int(ext.get("cx")) / cw if cw else 1.0)
                    gsy = sy * (int(ext.get("cy")) / chh if chh else 1.0)
                    child = (gx, gy, int(choff.get("x")), int(choff.get("y")), gsx, gsy)
                except Exception:
                    child = xf
                self.walk_shapes(shape.shapes, master, where, child, rects)
                continue
            entry = {"rect": rect, "framed": _is_framed(shape._element), "el": shape._element}
            if rect is not None:
                rects.append(entry)
            if getattr(shape, "has_table", False) and shape.has_table:
                self.add_table(shape, master, where, (sx, sy), rects, rect)
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                w = (rect[2] - rect[0]) if rect else None
                h = (rect[3] - rect[1]) if rect else None
                chain = self.chain_for_shape(shape, master)
                label = "titre" if chain.title else "zone de texte"
                if w and h:
                    label += " (%d x %d pt)" % (w, h)
                self.bodies.append(Body(shape.text_frame._txBody, chain, w, h, "shape",
                                        "%s, %s" % (where, label), shape=shape,
                                        slide_h=self.slide_h, rects=rects, rect=rect,
                                        framed=entry["framed"]))

    def add_table(self, gf, master, where, scale, rects, rect):
        tbl = gf.table
        cols = [c.width / EMU_PER_PT * scale[0] for c in tbl.columns]
        rows = [r.height / EMU_PER_PT * scale[1] for r in tbl.rows]
        theme = self.theme(master)
        for ri, row in enumerate(tbl.rows):
            for ci, cell in enumerate(row.cells):
                tc = cell._tc
                if tc.get("hMerge") in ("1", "true") or tc.get("vMerge") in ("1", "true"):
                    continue
                gs, rs = int(tc.get("gridSpan", "1")), int(tc.get("rowSpan", "1"))
                txb = tc.find(A + "txBody")
                if txb is None:
                    continue
                chain = Chain([txb.find(A + "lstStyle"), self.default_style],
                              [txb.find(A + "bodyPr")], theme)
                self.bodies.append(Body(txb, chain, sum(cols[ci:ci + gs]), sum(rows[ri:ri + rs]),
                                        "cell", "%s, cellule de tableau" % where, shape=gf,
                                        slide_h=self.slide_h, cell=(ri, ci, rs, tc),
                                        rects=rects, rect=rect, framed=True))

    def collect(self):
        prs = self.prs
        for n, slide in enumerate(prs.slides, start=1):
            master = slide.slide_layout.slide_master
            self.walk_shapes(slide.shapes, master, "diapositive %d" % n)
            if slide.has_notes_slide:
                tf = slide.notes_slide.notes_text_frame
                if tf is not None:
                    chain = Chain([tf._txBody.find(A + "lstStyle")], [], self.theme(master))
                    self.bodies.append(Body(tf._txBody, chain, kind="notes",
                                            context="diapositive %d, notes" % n))
        for master in prs.slide_masters:
            self.walk_shapes(master.shapes, master, "masque")
            for lay in master.slide_layouts:
                self.walk_shapes(lay.shapes, master, "disposition")
        for part in prs.part.package.iter_parts():
            ct = part.content_type
            if ct.endswith("drawingml.chart+xml") and hasattr(part, "_element"):
                for txb in part._element.iter(C + "rich"):
                    self.bodies.append(Body(txb, Chain([], [], None), kind="chart",
                                            context="titre de graphique"))
            elif ct.endswith("diagramData+xml") or ct.endswith("diagramDrawing+xml"):
                try:
                    root = etree.fromstring(part.blob)
                except Exception:
                    continue
                self.blob_parts.append((part, root))
                for el in root.iter():
                    if (el.tag.endswith("}t") or el.tag.endswith("}txBody")) and el.find(A + "p") is not None:
                        self.bodies.append(Body(el, Chain([], [], None), kind="smartart",
                                                context="SmartArt"))

def _is_framed(el):
    tag = el.tag.split("}")[-1]
    if tag in ("pic", "graphicFrame"):
        return True
    sppr = el.find(P + "spPr")
    style = el.find(P + "style")

    def ref_idx(name):
        if style is None:
            return 0
        r = style.find(A + name)
        return int(r.get("idx", "0")) if r is not None else 0

    fill = False
    if sppr is not None:
        if any(sppr.find(A + t) is not None for t in ("solidFill", "gradFill", "pattFill", "blipFill")):
            fill = True
        elif sppr.find(A + "noFill") is None and ref_idx("fillRef") > 0:
            fill = True
    line = False
    ln = sppr.find(A + "ln") if sppr is not None else None
    if ln is not None:
        if ln.find(A + "noFill") is None and (
                ln.find(A + "solidFill") is not None or ln.find(A + "gradFill") is not None
                or ref_idx("lnRef") > 0):
            line = True
    elif ref_idx("lnRef") > 0:
        line = True
    return fill or line

def _free_space(body):
    if body.rect is None:
        return 0.0, 0.0
    x0, y0, x1, y1 = body.rect
    below = (body.slide_h - y1 - 2.0) if body.slide_h else 0.0
    above = y0 - 2.0
    own = body.shape._element if body.shape is not None else None
    for r in body.rects:
        rr = r["rect"]
        if rr is None or r["el"] is own:
            continue
        if min(x1, rr[2]) - max(x0, rr[0]) <= 2.0:
            continue
        if rr[0] <= x0 + 1 and rr[2] >= x1 - 1 and rr[1] <= y0 + 1 and rr[3] >= y1 - 1:
            below = min(below, rr[3] - y1 - 3.0)
            above = min(above, y0 - rr[1] - 3.0)
        elif rr[1] >= y1 - 1:
            below = min(below, rr[1] - y1 - 2.0)
        elif rr[3] <= y0 + 1:
            above = min(above, y0 - rr[3] - 2.0)
        else:
            if rr[1] > y0:
                below = min(below, 0.0)
            if rr[3] < y1:
                above = min(above, 0.0)
    return max(0.0, below), max(0.0, above)

def _extra_room(body):
    if body.kind != "shape" or body.framed:
        return 0.0
    below, above = _free_space(body)
    anchor = body.chain.body_attr("anchor", "t")
    if anchor == "b":
        return above
    if anchor == "ctr":
        return 2 * min(below, above)
    return below

def chart_value_segments(prs):
    out = []
    for part in prs.part.package.iter_parts():
        if not part.content_type.endswith("drawingml.chart+xml") or not hasattr(part, "_element"):
            continue
        for cache in part._element.iter(C + "strCache"):
            for v in cache.iter(C + "v"):
                if v.text and v.text.strip():
                    seg = Segment()
                    seg.add_text(v.text, "-", None)
                    seg.meta["v"] = v
                    out.append(seg)
    return out

# --------------------------------------------------------------------------
# Ajustement
# --------------------------------------------------------------------------

WIDTH_SAFETY = 0.985

def _body_layout(body, para_runs, scale, width_factor=1.0):
    paras = [measure_paragraph(body.chain, p, runs) for p, runs in zip(body.paras, para_runs)
             if runs is not None]
    l, r, t, b = _insets(body)
    wrap = body.kind == "cell" or body.chain.body_attr("wrap", "square") != "none"
    width = (body.width or 0) - l - r
    h, lines, w = layout(paras, width * width_factor, scale, 0.0, wrap)
    return h, w, lines, (width, (body.height or 0) - t - b, wrap)

def _vertical(body):
    return body.chain.body_attr("vert", "horz") not in (None, "horz")

def needed_scale(body, orig_runs, new_runs, start_scale, extra=0.0):
    h_o, w_o, _, (avail_w, avail_h, wrap) = _body_layout(body, orig_runs, start_scale)
    limit_h = max(avail_h + extra, h_o) + 0.5
    limit_w = max(avail_w, w_o) + 0.5
    s = start_scale
    while s >= MIN_SCALE - 1e-9:
        h, w, _, _ = _body_layout(body, new_runs, s, WIDTH_SAFETY)
        if h <= limit_h and (wrap or w <= limit_w * WIDTH_SAFETY):
            return s
        s = round(s - SCALE_STEP, 4)
    return None

def overflow_at(body, orig_runs, new_runs, scale):
    h_o, _, _, (_, avail_h, _) = _body_layout(body, orig_runs, scale)
    h, _, _, _ = _body_layout(body, new_runs, scale, WIDTH_SAFETY)
    return max(0.0, h - max(avail_h, h_o) - 0.5)

def apply_scale(body, scale):
    own_bp = body.txBody.find(A + "bodyPr")
    if body.kind == "shape" and own_bp is not None and own_bp.find(A + "normAutofit") is not None:
        own_bp.find(A + "normAutofit").set("fontScale", str(int(round(scale * 100000))))
        return
    for p in body.paras:
        ppr = p.find(A + "pPr")
        lvl = int(ppr.get("lvl", "0")) if ppr is not None else 0
        for el in list(p):
            if el.tag in (A + "r", A + "fld", A + "br"):
                rpr = el.find(A + "rPr")
                if rpr is None:
                    rpr = etree.Element(A + "rPr")
                    el.insert(0, rpr)
            elif el.tag == A + "endParaRPr":
                rpr = el
            else:
                continue
            size, _ = run_metrics(body.chain, rpr, lvl)
            new = _round_size(size * scale)
            rpr.set("sz", str(int(round(new * 100))))

def _round_size(size):
    return max(1.0, math.floor(size * 2 + 1e-6) / 2)

def _runs_for(body, segments, results, orig):
    out = []
    for pi in range(len(body.paras)):
        i = body.seg_index.get(pi)
        if i is not None and results[i] is not None:
            out.append(translated_runs(segments[i], results[i]))
        elif pi in body.merged and _owner_translated(body, pi, results):
            out.append(None)
        else:
            out.append(orig[pi])
    return out

def _owner_translated(body, pi, results):
    k = pi
    while k >= 0 and k not in body.seg_index:
        k -= 1
    return k >= 0 and results[body.seg_index[k]] is not None

# --------------------------------------------------------------------------
# Fonction principale
# --------------------------------------------------------------------------

def translate_pptx_universal(input_path: str, output_path: str, glossary_path: str = None,
                             progress_callback=None, direction="Allemand ➔ Français"):
    start_time = time.time()
    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    work = tempfile.mkdtemp(prefix="traduction_")
    ext = os.path.splitext(input_path)[1]
    tmp_in = os.path.join(work, "in" + ext)
    tmp_out = os.path.join(work, "out" + ext)
    shutil.copy(input_path, tmp_in)

    tgt_lang = LANG[direction][1]
    prs = Presentation(tmp_in)
    glossary_df = load_glossary_df(glossary_path)

    inv = Inventory(prs)
    inv.collect()

    segments, contexts, owners = [], [], []
    for b in inv.bodies:
        pi = 0
        while pi < len(b.paras):
            group = [pi]
            while group[-1] + 1 < len(b.paras) and is_continuation(b.paras[group[-1]], b.paras[group[-1] + 1], b):
                group.append(group[-1] + 1)
            if len(group) > 1:
                seg = group_segment([b.paras[k] for k in group])
            else:
                seg = paragraph_segment(b.paras[pi])
            if seg is not None:
                b.seg_index[pi] = len(segments)
                for k in group[1:]:
                    b.merged.add(k)
                segments.append(seg)
                contexts.append(b.context)
                owners.append(b)
            pi = group[-1] + 1
            
    n_para = len(segments)
    chart_values = chart_value_segments(prs)
    segments += chart_values
    contexts += ["libellé de graphique"] * len(chart_values)

    orig = {id(b): [original_runs(p) for p in b.paras] for b in inv.bodies}

    # Échelonnement : translate_segments annonce 0 à 80 % (lots puis contrôle
    # qualité) avec son propre libellé ; ici la traduction couvre 0 à 70 %.
    def batch_progress(cur, tot, msg=None):
        if progress_callback:
            p = min(70, int((cur / max(1, tot)) * 100 * 70 / 80))
            progress_callback(p, 100, msg or "Traduction par l'IA...")

    results, stats = translate_segments(segments, direction, "pptx", glossary_df,
                                        batch_progress, contexts)

    if progress_callback:
        progress_callback(72, 100, "📐 Analyse géométrique de la mise en page...")

    report = {"shrunk": 0, "grown": 0, "shortened": 0, "overflow": 0, "table_grown": 0}
    plans = {}
    for b in inv.bodies:
        if b.kind not in ("shape", "cell") or not b.width or not b.height or _vertical(b):
            continue
        if not any(results[i] is not None for i in b.seg_index.values()):
            continue
        kind, el = b.chain.autofit()
        start = 1.0
        if kind == "normAutofit" and el is not None and el.get("fontScale"):
            start = int(el.get("fontScale")) / 100000.0
        try:
            extra = _extra_room(b)
            s = needed_scale(b, orig[id(b)], _runs_for(b, segments, results, orig[id(b)]),
                             start, extra)
        except Exception as exc:
            print(f"⚠️ Mesure impossible ({b.context}) : {exc}")
            continue
        plans[id(b)] = [b, start, s, extra]

    by_table = {}
    for plan in plans.values():
        if plan[0].kind == "cell":
            by_table.setdefault(id(plan[0].shape), []).append(plan)
    for cell_plans in by_table.values():
        if all(pl[2] is not None and pl[2] >= pl[1] - 1e-9 for pl in cell_plans):
            continue
        grow, grown_before = {}, {}
        for b, start, s, _ in cell_plans:
            new_runs = _runs_for(b, segments, results, orig[id(b)])
            row = b.cell[0]
            grow[row] = max(grow.get(row, 0.0), overflow_at(b, orig[id(b)], new_runs, start))
            h_o, _, _, (_, avail_h, _) = _body_layout(b, orig[id(b)], start)
            grown_before[row] = max(grown_before.get(row, 0.0), h_o - avail_h)
        below, _ = _free_space(cell_plans[0][0])
        if sum(grow.values()) <= below - sum(max(0.0, v) for v in grown_before.values()):
            for pl in cell_plans:
                pl[2] = pl[1]
            report["table_grown"] += 1

    to_shorten = []
    for b, start, s, extra in plans.values():
        if s is not None and s >= start * SHORTEN_BELOW:
            continue
        ratio = (s or MIN_SCALE) / start
        for i in b.seg_index.values():
            if results[i] is None:
                continue
            txt = items_to_tagged(segments[i], results[i])
            plain = re.sub(r"<[^>]+>", "", txt)
            if len(plain) > 12:
                to_shorten.append((i, segments[i], txt, max(6, int(len(plain) * ratio * ratio))))
                
    if to_shorten:
        if progress_callback:
            progress_callback(78, 100, f"✂️ Raccourcissement IA de {len(to_shorten)} textes trop longs...")
        for k in range(0, len(to_shorten), 40):
            short, i_tok, o_tok = shorten_segments(to_shorten[k:k + 40], direction)
            stats["input_tokens"] += i_tok
            stats["output_tokens"] += o_tok
            for i, items in short.items():
                b = owners[i]
                plan = plans.get(id(b))
                if plan is None:
                    continue
                old = results[i]
                results[i] = items
                s2 = needed_scale(b, orig[id(b)], _runs_for(b, segments, results, orig[id(b)]),
                                  plan[1], plan[3])
                if (s2 or 0) + 1e-9 > (plan[2] or 0):
                    plan[2] = s2
                    report["shortened"] += 1
                else:
                    results[i] = old

    if progress_callback:
        progress_callback(88, 100, "✍️ Réécriture des formes et styles XML...")

    for seg, items in zip(segments[:n_para], results[:n_para]):
        if items is not None:
            rebuild_paragraph(seg, items, tgt_lang)
    for seg, items in zip(segments[n_para:], results[n_para:]):
        if items is not None:
            seg.meta["v"].text = "".join(it[2] for it in items if it[0] == "t")
    for part, root in inv.blob_parts:
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    for b, start, s, extra in plans.values():
        if s is not None and s >= start - 1e-9:
            continue
        kind, _ = b.chain.autofit()
        if b.kind == "shape" and kind == "spAutoFit":
            h_new, _, _, (_, avail_h, _) = _body_layout(
                b, _runs_for(b, segments, results, orig[id(b)]), start, WIDTH_SAFETY)
            need = h_new - avail_h + 1.0
            below, _ = _free_space(b)
            if 0 < need <= below:
                b.shape.height = int(b.shape.height + need * EMU_PER_PT)
                report["grown"] += 1
                continue
        if s is None:
            report["overflow"] += 1
        final = s if s is not None else MIN_SCALE
        own_bp = b.txBody.find(A + "bodyPr")
        if own_bp is not None and own_bp.find(A + "normAutofit") is not None:
            apply_scale(b, final)
        else:
            apply_scale(b, final / start)
        report["shrunk"] += 1

    if progress_callback:
        progress_callback(96, 100, "💾 Enregistrement du fichier PowerPoint final...")

    prs.save(tmp_out)
    shutil.copy(tmp_out, output_path)
    shutil.rmtree(work, ignore_errors=True)

    if progress_callback:
        progress_callback(100, 100, "✅ Traduction finalisée !")

    print(f"📐 Ajustement : {report['shrunk']} cadres réduits, {report['grown']} agrandis, "
          f"{report['table_grown']} tableaux agrandis, "
          f"{report['shortened']} textes raccourcis, {report['overflow']} encore trop longs")
    print(f"🔤 Segments : {stats['segments']}, envoyés : {stats['segments'] - stats['skipped']}, "
          f"redemandés : {stats['retried']}, sans styles locaux : {stats['fallback']}")
    return {
        "status": "success",
        "output_path": output_path,
        "time_seconds": round(time.time() - start_time, 2),
        "input_tokens": stats["input_tokens"],
        "output_tokens": stats["output_tokens"],
        "total_cost_usd": stats["input_tokens"] / 1e6 * PRICE_IN + stats["output_tokens"] / 1e6 * PRICE_OUT,
        "fit": report,
    }