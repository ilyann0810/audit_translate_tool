# -*- coding: utf-8 -*-
"""
Genere des documents de test synthetiques dans tests/samples/.

Chaque fichier concentre des pieges connus, pour verifier que l'outil
reste general (et pas seulement regle sur les documents reels) :

  xlsx : texte riche (mot en gras rouge), formule NB.SI dont la cle texte
         ne doit pas etre traduite, commentaire, graphique avec titre,
         image, cellules fusionnees, deux feuilles.
  pptx : groupe, tableau, lien hypertexte, mot en couleur, notes, cadre en
         "reduire le texte", cadre en "ajuster la forme", graphique,
         document mixte francais / allemand.
  docx : image au milieu d'un paragraphe, zone de texte, lien hypertexte,
         champ de page dans le pied de page, surlignage et couleur, tableau,
         en-tete.

    python tests/make_samples.py
"""
import io
import os

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "samples")


def png_bytes(color=(200, 30, 30)):
    import fitz
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 60, 40), False)
    pix.set_rect(pix.irect, color)
    return pix.tobytes("png")


def make_xlsx(path):
    import openpyxl
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daten"
    ws["A1"] = "Filiale"
    ws["B1"] = "Umsatz (Mio. €)"
    ws["C1"] = "Kontrolle durchgeführt"
    rows = [("Filiale 4440 Nemours", 12.5, "Ja"), ("Filiale 4451 Toulon", 9.8, "Nein"),
            ("Filiale 4453 Saint Raphaël", 11.2, "Ja"), ("Filiale 4194 Erquy", 7.4, "Ja")]
    for i, (a, b, c) in enumerate(rows, start=2):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
    ws["A7"] = "Anzahl durchgeführter Kontrollen"
    ws["C7"] = '=COUNTIF(C2:C5,"Ja")'
    ws["A9"] = CellRichText([TextBlock(InlineFont(b=True, color="FFFF0000"), "Achtung: "),
                             "Die Werte der letzten Woche sind noch nicht vollständig."])
    ws.merge_cells("A11:C11")
    ws["A11"] = "Zusammenfassung der Prüfungsergebnisse für das Management"
    ws["A1"].comment = Comment("Die Filialnummer muss immer angegeben werden.", "Revision")
    chart = BarChart()
    chart.title = "Umsatz pro Filiale"
    chart.y_axis.title = "Umsatz"
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=5))
    ws.add_chart(chart, "E2")
    ws.add_image(Image(io.BytesIO(png_bytes())), "E20")
    ws2 = wb.create_sheet("Hinweise")
    ws2["A1"] = "Diese Tabelle wird jede Woche aktualisiert."
    ws2["A2"] = "Verantwortlich: L_FR_Logistik"
    wb.save(path)


def make_pptx(path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Pt, Inches

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Résultats de l'audit des supermarchés"
    tb = s.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r1 = p.add_run()
    r1.text = "Le contrôle des températures n'a pas été "
    r2 = p.add_run()
    r2.text = "effectué"
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    r3 = p.add_run()
    r3.text = " dans 4 supermarchés sur 10."
    for r in (r1, r2, r3):
        r.font.size = Pt(14)
    p2 = tf.add_paragraph()
    r = p2.add_run()
    r.text = "Plus d'informations sur l'intranet"
    r.font.size = Pt(12)
    r.hyperlink.address = "https://example.com/audit"
    tf.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT

    box = s.shapes.add_textbox(Inches(5), Inches(1.5), Inches(4), Inches(1.2))
    box.text_frame.word_wrap = True
    box.text_frame.text = "Wochenkennzahlen | Mise à jour en cours selon le CuR final"
    box.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    box.line.color.rgb = RGBColor(0, 0, 0)

    grp = s.shapes.add_group_shape()
    g1 = grp.shapes.add_textbox(Inches(0.5), Inches(3), Inches(2), Inches(0.5))
    g1.text_frame.text = "Responsable : Direction régionale"
    g2 = grp.shapes.add_textbox(Inches(2.6), Inches(3), Inches(2), Inches(0.5))
    g2.text_frame.text = "Échéance : fin du trimestre"

    rows, cols = 3, 2
    t = s.shapes.add_table(rows, cols, Inches(0.5), Inches(4), Inches(5), Inches(1.2)).table
    data = [("Thème", "Constat"), ("Signalétique", "Non conforme dans 3 supermarchés"),
            ("Éclairage", "Zones prédéfinies non respectées")]
    for i, (a, b) in enumerate(data):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b

    s.notes_slide.notes_text_frame.text = "Présenter d'abord les résultats globaux, puis les détails."

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Évolution du chiffre d'affaires"
    cd = CategoryChartData()
    cd.categories = ["Frais", "Surgelés", "Épicerie"]
    cd.add_series("Semaine 41", (12.1, 8.4, 15.0))
    ch = s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5),
                             Inches(6), Inches(4), cd).chart
    ch.has_title = True
    ch.chart_title.text_frame.text = "Chiffre d'affaires par rayon"
    prs.save(path)


def make_docx(path):
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml import parse_xml
    from docx.shared import Pt, RGBColor, Inches

    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Interne Revision – Prüfbericht"
    fp = sec.footer.paragraphs[0]
    fp.add_run("Seite ")
    fld = parse_xml('<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                    'w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple>')
    fp._p.append(fld)

    d.add_heading("Prüfung der Kühlkette", level=1)
    p = d.add_paragraph("Die Kerntemperatur wurde ")
    r = p.add_run("nicht")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    p.add_run(" bei allen Lieferungen gemessen. ")
    r = p.add_run("Dringender Handlungsbedarf")
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    p.add_run(".")

    # Image au milieu du texte
    img = io.BytesIO(png_bytes((30, 120, 200)))
    p = d.add_paragraph("Das folgende Symbol ")
    p.add_run().add_picture(img, width=Inches(0.3))
    p.add_run(" kennzeichnet kritische Feststellungen.")

    # Lien hypertexte
    p = d.add_paragraph("Weitere Informationen finden Sie im ")
    rid = d.part.relate_to("https://example.com/handbuch",
                           "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    p._p.append(parse_xml(
        '<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="%s">'
        '<w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr><w:t>Revisionshandbuch</w:t></w:r>'
        '</w:hyperlink>' % rid))
    p.add_run(".")

    # Zone de texte (VML simple, lue par Word)
    p = d.add_paragraph()
    p._p.append(parse_xml(
        '<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        'xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>'
        '<v:shape style="width:250pt;height:40pt" stroked="t"><v:textbox><w:txbxContent>'
        '<w:p><w:r><w:t>Hinweis: Die Ergebnisse sind vertraulich.</w:t></w:r></w:p>'
        '</w:txbxContent></v:textbox></v:shape></w:pict></w:r>'))

    t = d.add_table(rows=3, cols=2)
    t.style = "Table Grid"
    for i, (a, b) in enumerate([("Prüfthema", "Feststellung"),
                                ("Beschilderung", "Nicht konform in 3 Filialen"),
                                ("Beleuchtung", "Vordefinierte Zonen nicht eingehalten")]):
        t.cell(i, 0).text = a
        t.cell(i, 1).text = b
    d.add_paragraph("Verantwortlich: L_FR_Logistik (MID: 1596225)")
    d.save(path)


def main():
    os.makedirs(OUT, exist_ok=True)
    make_xlsx(os.path.join(OUT, "synthetique_de.xlsx"))
    make_pptx(os.path.join(OUT, "synthetique_fr_mixte.pptx"))
    make_docx(os.path.join(OUT, "synthetique_de.docx"))
    print("ok :", sorted(os.listdir(OUT)))


if __name__ == "__main__":
    main()
